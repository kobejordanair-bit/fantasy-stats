"""CSV, typed XLSX and standalone offline HTML from one validated canonical season."""
from __future__ import annotations
import csv
import html
import json
import math
import unicodedata
from datetime import datetime
from pathlib import Path
from fantasy_core import PAIRS, awards, validate_season

SHEETS=['球員季總','球員週均','球隊週次','球員逐週','對戰結果','類別戰績','交易紀錄','補人紀錄','交易比較','補人貢獻']
CSV_NAMES=['season_summary','season_weekly_average','team_weekly_summary','player_weekly_detail',
           'matchup_results','category_record','trade_history','waiver_history','trade_roi','waiver_roi']


def source_note(season):
    scopes=sorted({r.get('scope','unknown') for r in season['playerWeeks']})
    note='來源：'+season.get('source',{}).get('kind','unknown')+'。球隊口徑：'+season.get('quality',{}).get('scope','unknown')+'。球員口徑：'+(', '.join(scopes) or '未記錄')+'。空白表示未知。'
    warnings=season.get('quality',{}).get('warnings',[])
    if warnings: note+=' 提醒：'+('；'.join(warnings)[:500])+'（完整提醒見 season.json）'
    return note


def source_note_height(note, column_widths, *, minimum=34, font_size=11):
    """Estimate wrapped text height for the merged source row, in Excel points.

    Excel does not auto-fit merged cells. Column widths are approximately seven
    pixels per Calibri 11 digit plus five pixels padding. Reserve inner padding,
    count CJK/full-width glyphs as one em, and allow an extra line for long notes.
    This changes presentation only; the complete source-note string is retained.
    """
    available=max(1,sum(width*7+5 for width in column_widths)-12)
    em=font_size*96/72
    lines=0
    for paragraph in note.replace('\r\n','\n').replace('\r','\n').split('\n'):
        width=0
        for char in paragraph:
            if unicodedata.combining(char) or unicodedata.category(char)=='Cf': continue
            width+=em*(2 if char=='\t' else 1 if unicodedata.east_asian_width(char) in {'W','F','A'} else .52)
        lines+=max(1,math.ceil(width/available))
    return min(409,max(minimum,(lines+(1 if lines>2 else 0))*font_size*1.3+8))


def safe_csv(value):
    if isinstance(value,str) and (value.lstrip().startswith(('=','+','-','@')) or value.startswith(('\t','\r','\n'))):
        return "'"+value
    return value


def report_tables(s):
    keys=[c['key'] for c in s['categories']]
    auxiliary=[k for pair in PAIRS.values() for k in pair if k not in keys]
    stat_keys=keys+list(dict.fromkeys(auxiliary))
    tables=[]
    players=sorted(s['players'],key=lambda p:(p['stats'].get('PTS') is not None,p['stats'].get('PTS') or 0),reverse=True)
    base=['球員代碼','球員','位置','狀態','在陣週數']
    tables.append((base+stat_keys,[[p['id'],p['name'],p['position'],p['status'],p['weeks']]+[p['stats'].get(k) for k in stat_keys] for p in players]))
    tables.append((base+['週均 '+k if '%' not in k else k for k in stat_keys],[[p['id'],p['name'],p['position'],p['status'],p['weeks']]+[
        (p['stats'].get(k) if '%' in k else p['stats'][k]/p['weeks']) if p['stats'].get(k) is not None and p['weeks'] else None for k in stat_keys] for p in players]))
    tables.append((['週次','開始日','結束日','狀態','階段']+stat_keys,[[r['week'],r['startDate'],r['endDate'],r['status'],r['phase']]+[r['stats'].get(k) for k in stat_keys] for r in s['teamWeeks']]))
    tables.append((['週次','球員代碼','球員','位置','統計口徑']+stat_keys,[[r['week'],r['playerId'],r['name'],r['position'],r['scope']]+[r['stats'].get(k) for k in stat_keys] for r in s['playerWeeks']]))
    tables.append((['週次','對手','狀態','階段','結果','類別勝','類別敗','類別平'],[[r.get(k) for k in ['week','opponent','status','phase','result','categoryWins','categoryLosses','categoryTies']] for r in s['matchups']]))
    tables.append((['類別','勝','敗','平','勝率%'],[[r['key'],r['wins'],r['losses'],r['ties'],r['wins']/(r['wins']+r['losses']+r['ties']) if r['wins']+r['losses']+r['ties'] else None] for r in s['categoryRecord']]))
    refs=lambda people:', '.join(p['name']+' ['+p['id']+']' for p in people)
    tables.append((['交易代碼','時間','交易對手','我方收到','我方送出','筆記'],[[r['id'],r['date'],r['opponent'],refs(r['received']),refs(r['sent']),r['notes']] for r in s['transactions'] if r['type']=='trade']))
    tables.append((['異動代碼','時間','類型','取得','釋出','筆記'],[[r['id'],r['date'],r['type'],refs(r['received']),refs(r['sent']),r['notes']] for r in s['transactions'] if r['type']!='trade']))
    trade_headers=['日期','交易對手','我方收到','我方送出','從第N週起算','計算口徑']+[prefix+k for k in keys for prefix in ['我方_','對方_','結果_']]
    trades=[]
    for r in s['tradeAnalysis']:
        cols={c['key']:c for c in r['categories']}
        trades.append([r[k] for k in ['date','opponent','received','sent','fromWeek','method']]+[cols.get(k,{}).get(v) for k in keys for v in ['received','sent','result']])
    tables.append((trade_headers,trades))
    tables.append((['撿入日期','球員','從第N週','在陣週數','計算口徑']+['累積_'+k for k in keys]+['週均_'+k for k in keys],
                   [[r[k] for k in ['date','name','fromWeek','weeks','method']]+[r['stats'].get(k) for k in keys]+[r['weeklyAverage'].get(k) for k in keys] for r in s['waiverAnalysis']]))
    return tables


def export_csvs(season,out_dir):
    for name,(headers,rows) in zip(CSV_NAMES,report_tables(season)):
        with (Path(out_dir)/(name+'.csv')).open('x',encoding='utf-8-sig',newline='') as stream:
            writer=csv.writer(stream)
            writer.writerow([safe_csv(v) for v in headers])
            writer.writerows([[safe_csv(v) for v in row] for row in rows])


def export_excel(season,out_dir):
    # Python library is intentionally a runtime dependency of the reusable collector.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    workbook=Workbook();workbook.remove(workbook.active)
    for name,(headers,rows) in zip(SHEETS,report_tables(season)):
        ws=workbook.create_sheet(name)
        ws.append([season['season']+' '+season['teamName']+' — '+name])
        ws.append([source_note(season)])
        ws.append(headers)
        for row in rows: ws.append(row)
        for row in ws:
            for cell in row:
                if isinstance(cell.value,str): cell.data_type='s'  # Including strings starting with =.
                cell.font=Font(name='Calibri',size=11,color='18283D')
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for row_no in [1,3]:
            for cell in ws[row_no]:
                cell.fill=PatternFill('solid',fgColor='142A43')
                cell.font=Font(name='Calibri',bold=True,color='FFFFFF',size=14 if row_no==1 else 11)
        for col,header in enumerate(headers,1):
            letter=get_column_letter(col)
            wide=header in {'球員','交易對手','我方收到','我方送出','取得','釋出','筆記','計算口徑'}
            ws.column_dimensions[letter].width=34 if wide else 20 if '日' in header or '時間' in header else 15
            for r in range(4,ws.max_row+1):
                cell=ws.cell(r,col)
                if type(cell.value) in {float,int}:
                    cell.number_format='0.0%' if '%' in header else '#,##0.00' if '週均' in header else '#,##0.##'
                if header in {'開始日','結束日','撿入日期','日期','時間'} and isinstance(cell.value,str) and cell.value:
                    try:
                        # Dates retain source clock text if they include time/zone; never strip a timestamp.
                        if len(cell.value)==10:
                            cell.value=datetime.strptime(cell.value,'%Y-%m-%d');cell.number_format='yyyy-mm-dd'
                    except ValueError: pass
        ws.row_dimensions[1].height=28
        ws.row_dimensions[2].height=source_note_height(
            ws['A2'].value,[ws.column_dimensions[get_column_letter(col)].width for col in range(1,len(headers)+1)],
            minimum=64 if season.get('quality',{}).get('warnings') else 34)
        ws.row_dimensions[3].height=32
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers))
        ws.freeze_panes='F4';ws.auto_filter.ref=f'A3:{get_column_letter(len(headers))}{max(3,ws.max_row)}'
        ws.sheet_view.showGridLines=False;ws.print_title_rows='1:3'
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    target=Path(out_dir)/'fantasy_season.xlsx'
    if target.exists(): raise FileExistsError(target)
    workbook.save(target);workbook.close()
    return target


def export_html(season,out_dir):
    leaders,mvp=awards(season['players'],season['categories'])
    payload={'season':season,'tables':[{'name':name,'headers':headers,'rows':rows} for name,(headers,rows) in zip(SHEETS,report_tables(season))],
             'leaders':leaders,'mvp':mvp,'sourceNote':source_note(season)}
    # Escape every literal <, not only lowercase </script>. Render data solely with textContent.
    encoded=json.dumps(payload,ensure_ascii=False,allow_nan=False).replace('<','\\u003c').replace('>','\\u003e').replace('&','\\u0026').replace('\u2028','\\u2028').replace('\u2029','\\u2029')
    template='''<!doctype html><html lang="zh-Hant"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; script-src 'unsafe-inline'; style-src 'unsafe-inline'; img-src data:; connect-src 'none'; base-uri 'none'; form-action 'none'">
<title>__TITLE__</title><style>body{margin:0;background:#eff3f7;color:#18283d;font:16px system-ui}main{max-width:1500px;margin:auto;padding:24px}h1{font-size:28px}nav{display:flex;gap:8px;flex-wrap:wrap;margin:20px 0}button,select{font:inherit;border:1px solid #aab8c8;border-radius:5px;padding:8px;background:white;color:#18283d}button[aria-pressed=true]{background:#142a43;color:white}.scroll{overflow:auto;background:white;border:1px solid #ccd6e2}table{border-collapse:collapse;width:100%;font-size:14px}th,td{padding:10px;text-align:left;border-bottom:1px solid #dfe5ee;white-space:nowrap}th{background:#142a43;color:white;cursor:pointer;position:sticky;top:0}caption{text-align:left;padding:10px}.note{color:#455d74;line-height:1.7}.chart{background:white;padding:12px;margin-bottom:15px}.bar{background:#196b97;height:18px;display:inline-block}.warning{border-left:3px solid #a86700;padding-left:14px}@media(max-width:600px){main{padding:12px}h1{font-size:23px}}</style>
<main><h1>__TITLE__</h1><p class="note" id="intro"></p><div id="warnings" class="note warning"></div><nav id="nav" aria-label="報表"></nav><section id="content"></section></main>
<script id="data" type="application/json">__DATA__</script><script>
'use strict';const D=JSON.parse(document.getElementById('data').textContent),root=document.getElementById('content');
const node=(tag,text)=>{const n=document.createElement(tag);if(text!==undefined)n.textContent=String(text);return n};
const fmt=(v,h)=>v===null||v===undefined?'—':typeof v==='number'?(h.includes('%')?(v*100).toFixed(1)+'%':v.toLocaleString('zh-TW',{maximumFractionDigits:3})):String(v);
document.getElementById('intro').textContent=D.sourceNote+' 所有內容均可離線使用。';
for(const text of D.season.quality.warnings)document.getElementById('warnings').append(node('p',text));
function table(headers,rows){const wrap=node('div');wrap.className='scroll';const t=node('table'),head=node('thead'),body=node('tbody');t.append(head,body);wrap.append(t);let order=1;function paint(data){body.replaceChildren();for(const r of data){const tr=node('tr');r.forEach((v,i)=>tr.append(node('td',fmt(v,headers[i]))));body.append(tr)}if(!data.length){const td=node('td','尚無可用紀錄');td.colSpan=headers.length;body.append(td)}}const hr=node('tr');headers.forEach((h,i)=>{const th=node('th',h);th.tabIndex=0;const sort=()=>{order*=-1;paint([...rows].sort((a,b)=>a[i]==null?1:b[i]==null?-1:typeof a[i]==='number'&&typeof b[i]==='number'?order*(a[i]-b[i]):order*String(a[i]).localeCompare(String(b[i]))))};th.onclick=sort;th.onkeydown=e=>{if(e.key==='Enter')sort()};hr.append(th)});head.append(hr);paint(rows);return wrap}
function show(name){
 root.replaceChildren();for(const b of document.querySelectorAll('nav button'))b.setAttribute('aria-pressed',b.textContent===name);
 root.append(node('h2',name));const report=D.tables.find(t=>t.name===name);
 if(report){root.append(table(report.headers,report.rows));return}
 if(name==='MVP'){root.append(node('p','MVP 為隊內累積計數類別 Z-score 合計，負向類別反轉；命中率不納入。任何類別缺漏則不給總分，不等同聯盟球員排名。'));const keys=D.season.categories.filter(c=>c.kind==='count').map(c=>c.key);root.append(table(['球員','週數','綜合 Z-score',...keys],D.mvp.map(p=>[p.name,p.weeks,p.score,...keys.map(k=>p.categories[k])])));return}
 if(name==='類別領先者'){root.append(node('p','負向類別列最多失誤／犯規者，並非價值最佳；比例保留整段命中率，未設最低出手門檻，請同時查看出手數。'));root.append(table(['類別','口徑','累積球員','累積','週均球員','週均／比例'],D.leaders.map(r=>[r.category,r.label,r.totalName,r.total,r.averageName,fmt(r.average,r.category)])));return}
 const select=node('select');select.setAttribute('aria-label','走勢類別');for(const c of D.season.categories)select.append(node('option',c.key));
 const playerSelect=node('select');playerSelect.setAttribute('aria-label','走勢球員');
 const ids=new Map(D.season.playerWeeks.map(r=>[r.playerId||r.name,r.name]));
 for(const [id,label] of ids){const option=node('option',label);option.value=id;playerSelect.append(option)}
 const playerMode=name==='球員走勢',area=node('div');root.append(select);if(playerMode)root.append(playerSelect);root.append(area);
 const draw=()=>{
  area.replaceChildren();const key=select.value;
  const rows=playerMode?D.season.playerWeeks.filter(r=>(r.playerId||r.name)===playerSelect.value):D.season.teamWeeks;
  if(!rows.length){area.append(node('p',playerMode?'沒有可用的球員逐週實際資料。':'沒有可用的球隊週次資料。'));return}
  const max=Math.max(1,...rows.map(r=>r.stats[key]??0));
  for(const r of rows){const p=node('p','W'+r.week+' '+(playerMode?r.scope:r.status)+' '+fmt(r.stats[key],key)+' ');if(r.stats[key]!=null){const bar=node('span');bar.className='bar';bar.style.width=Math.max(1,r.stats[key]/max*65)+'%';p.append(bar)}area.append(p)}
 };
 select.onchange=draw;playerSelect.onchange=draw;draw()
}
for(const name of ['球隊走勢','球員走勢',...D.tables.map(t=>t.name),'類別領先者','MVP']){const b=node('button',name);b.onclick=()=>show(name);document.getElementById('nav').append(b)}show('球隊走勢');
</script></html>'''
    title=html.escape(season['season']+' '+season['teamName']+' 年度戰情室',quote=True)
    result=template.replace('__TITLE__',title).replace('__DATA__',encoded)
    target=Path(out_dir)/'fantasy_dashboard.html';target.write_text(result,encoding='utf-8')
    return target


def export_all(season,out_dir):
    validate_season(season);out_dir=Path(out_dir);out_dir.mkdir(parents=True,exist_ok=True)
    targets=[out_dir/'season.json',out_dir/'fantasy_dashboard.html',out_dir/'fantasy_season.xlsx']+[out_dir/(n+'.csv') for n in CSV_NAMES]
    if any(p.exists() for p in targets): raise FileExistsError('匯出目錄已有報表；請選擇新目錄以保留舊資料。')
    (out_dir/'season.json').write_text(json.dumps(season,ensure_ascii=False,allow_nan=False,indent=2),encoding='utf-8')
    export_csvs(season,out_dir);export_excel(season,out_dir);export_html(season,out_dir)
    return out_dir
