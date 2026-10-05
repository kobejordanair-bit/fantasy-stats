import csv
import json
import sys
import tempfile
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from fantasy_core import blank_season
from season_exports import export_all, safe_csv, SHEETS
from season_exports import source_note
from fetch_stats import new_run


def fixture():
    s=blank_season(2025,'=SYNTHETIC </ScRiPt><img src=x onerror=alert(1)>')
    s['categories']=[{'key':'PTS','name':'PTS','kind':'count','lowerIsBetter':False},{'key':'FG%','name':'FG%','kind':'percentage','lowerIsBetter':False}]
    s['players']=[{'id':'999.p.1','name':'=SUM(1,2)','position':'PG','status':'observed','weeks':2,'stats':{'PTS':10,'FG%':.5,'FGM':5,'FGA':10}}]
    s['quality']['warnings']=['Synthetic test only; no real league data.']
    return s


class ExportTests(unittest.TestCase):
    def test_typed_excel_csv_escape_html_offline_and_non_overwrite(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            export_all(fixture(),tmp)
            with open(Path(tmp)/'season_summary.csv',encoding='utf-8-sig',newline='') as f:
                rows=list(csv.reader(f))
            self.assertEqual(rows[1][1],"'=SUM(1,2)")
            self.assertEqual(float(rows[1][rows[0].index('FG%')]),.5)
            book=load_workbook(Path(tmp)/'fantasy_season.xlsx',data_only=False)
            self.assertEqual(book.sheetnames,SHEETS)
            ws=book[SHEETS[0]];headers=[cell.value for cell in ws[3]]
            cell=ws.cell(4,headers.index('FG%')+1)
            self.assertEqual(cell.value,.5);self.assertEqual(cell.data_type,'n');self.assertEqual(cell.number_format,'0.0%')
            self.assertEqual(ws['B4'].data_type,'s');self.assertEqual(ws['B4'].value,'=SUM(1,2)')
            self.assertEqual(ws.freeze_panes,'F4');self.assertIsNotNone(ws.auto_filter.ref);book.close()
            markup=(Path(tmp)/'fantasy_dashboard.html').read_text(encoding='utf-8')
            self.assertNotIn('</ScRiPt>',markup);self.assertNotIn('<img src=x',markup)
            self.assertNotIn('<script src=',markup);self.assertIn('\\u003c',markup)
            self.assertNotIn('innerHTML',markup);self.assertIn('類別領先者',markup);self.assertIn('MVP',markup)
            self.assertIn('球員走勢',markup);self.assertIn('走勢球員',markup)
            self.assertEqual(json.loads((Path(tmp)/'season.json').read_text(encoding='utf-8'))['players'][0]['stats']['FG%'],.5)
            with self.assertRaises(FileExistsError):export_all(fixture(),tmp)

    def test_csv_formula_prefixes_and_zero(self):
        for value in ['=1',' +SUM(1)','-formula','@SUM(1)','\tcmd','\rtext','\ntext']:
            self.assertTrue(safe_csv(value).startswith("'"))
        self.assertEqual(safe_csv(0),0);self.assertEqual(safe_csv(-2),-2)

    def test_annual_runs_never_reuse_previous_folder(self):
        with tempfile.TemporaryDirectory() as tmp:
            first=new_run(tmp,'2025-26');second=new_run(tmp,'2025-26')
            self.assertNotEqual(first,second);self.assertTrue(first.exists());self.assertTrue(second.exists())

    def test_offline_manual_source_is_not_promoted_to_official(self):
        s=fixture();s['source']['kind']='manual';s['quality']['scope']='roster-week'
        note=source_note(s)
        self.assertIn('manual',note);self.assertIn('roster-week',note)
        self.assertNotIn('官方',note);self.assertNotIn('已驗證',note)

if __name__=='__main__':unittest.main()
